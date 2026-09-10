"""
Document ingestion dry run (Bedrock migration plan #6).

Parses contract PDFs/DOCX with open-source parsers (pypdf, python-docx, and
`unstructured` when installed) and reports on STRUCTURE — the question this
answers is: will Bedrock Knowledge Bases' default parser suffice, or do the
documents have layouts (nested tables, multi-column pages, scanned images)
that need the FM-based / Bedrock Data Automation parser?

No AWS dependency; runs fully offline.

Usage:
  python doc_ingestion/generate_samples.py          # once, to create samples/
  python doc_ingestion/dry_run.py                   # parse ./doc_ingestion/samples
  python doc_ingestion/dry_run.py --dir path/to/real/contracts
"""

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

_CLAUSE_RE = re.compile(
    r"(force majeure|mitigation|liabilit|indemnif|termination|klausul|"
    r"keadaan kahar)", re.IGNORECASE)


# ── DOCX ─────────────────────────────────────────────────────────────────────

def _walk_docx_tables(table, depth=1):
    """Yield (depth, n_rows, n_cols) for a table and any tables nested in cells."""
    yield depth, len(table.rows), len(table.columns)
    for row in table.rows:
        for cell in row.cells:
            for nested in cell.tables:
                yield from _walk_docx_tables(nested, depth + 1)


def inspect_docx(path: Path) -> dict:
    from docx import Document
    doc = Document(str(path))
    headings = [p.text for p in doc.paragraphs
                if p.style.name.startswith("Heading") and p.text.strip()]
    body_chars = sum(len(p.text) for p in doc.paragraphs)
    tables = []
    for t in doc.tables:
        tables.extend(_walk_docx_tables(t))
    clause_hits = sorted({m.group(1).lower()
                          for p in doc.paragraphs
                          for m in [_CLAUSE_RE.search(p.text)] if m})
    return {
        "file": path.name,
        "format": "docx",
        "paragraphs": len(doc.paragraphs),
        "text_chars": body_chars,
        "headings": len(headings),
        "heading_samples": headings[:5],
        "tables_total": len(tables),
        "nested_tables": sum(1 for d, _, _ in tables if d > 1),
        "table_shapes": [f"depth{d}:{r}x{c}" for d, r, c in tables],
        "clause_keywords_found": clause_hits,
        "multi_column_suspected": False,  # DOCX columns live in section props; rare here
        "notes": [],
    }


# ── PDF ──────────────────────────────────────────────────────────────────────

def _page_column_signal(page) -> bool:
    """Heuristic multi-column detector: collect x-positions of text fragments
    and look for two dense clusters separated by a wide gap near mid-page."""
    xs: list[float] = []

    def visitor(text, cm, tm, font_dict, font_size):
        if text and text.strip():
            xs.append(tm[4])

    try:
        page.extract_text(visitor_text=visitor)
    except Exception:
        return False
    if len(xs) < 30:
        return False
    width = float(page.mediabox.width)
    left = sum(1 for x in xs if x < width * 0.40)
    right = sum(1 for x in xs if x >= width * 0.48)
    mid = sum(1 for x in xs if width * 0.40 <= x < width * 0.48)
    return left >= 10 and right >= 10 and mid <= min(left, right) * 0.2


def inspect_pdf(path: Path) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    n_pages = len(reader.pages)
    texts, col_pages, image_pages = [], [], 0
    for i, page in enumerate(reader.pages):
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
        if _page_column_signal(page):
            col_pages.append(i + 1)
        if page.images:
            image_pages += 1
    all_text = "\n".join(texts)
    clause_hits = sorted({m.group(1).lower() for m in _CLAUSE_RE.finditer(all_text)})
    empty_pages = sum(1 for t in texts if len(t.strip()) < 20)
    notes = []
    if empty_pages:
        notes.append(f"{empty_pages}/{n_pages} pages have almost no extractable "
                     "text — likely scanned images; would need OCR (FM parser / "
                     "Bedrock Data Automation, or local Textract-equivalent).")
    if col_pages:
        notes.append(f"multi-column layout suspected on pages {col_pages} — "
                     "plain text extraction interleaves columns; chunking by "
                     "layout region needed.")
    return {
        "file": path.name,
        "format": "pdf",
        "pages": n_pages,
        "text_chars": len(all_text),
        "empty_text_pages": empty_pages,
        "pages_with_images": image_pages,
        "multi_column_pages": col_pages,
        "multi_column_suspected": bool(col_pages),
        "clause_keywords_found": clause_hits,
        "notes": notes,
    }


# ── optional: unstructured (element-level structure) ─────────────────────────

def inspect_with_unstructured(path: Path) -> dict | None:
    try:
        from unstructured.partition.auto import partition
    except ImportError:
        return None
    try:
        elements = partition(filename=str(path))
    except Exception as exc:
        return {"error": str(exc)[:200]}
    kinds: dict[str, int] = {}
    for el in elements:
        kinds[type(el).__name__] = kinds.get(type(el).__name__, 0) + 1
    return {"element_counts": kinds}


# ── report ───────────────────────────────────────────────────────────────────

def run(sample_dir: Path) -> list[dict]:
    reports = []
    files = sorted([*sample_dir.glob("*.pdf"), *sample_dir.glob("*.docx")])
    if not files:
        print(f"No PDF/DOCX files in {sample_dir} — run generate_samples.py "
              "or pass --dir pointing at real contracts.")
        return reports
    for f in files:
        rep = inspect_pdf(f) if f.suffix == ".pdf" else inspect_docx(f)
        extra = inspect_with_unstructured(f)
        if extra:
            rep["unstructured"] = extra
        reports.append(rep)
        print(json.dumps(rep, indent=1))
    out = sample_dir.parent / "dry_run_report.json"
    out.write_text(json.dumps(
        {"generated_at": datetime.now().isoformat(timespec="seconds"),
         "source_dir": str(sample_dir), "documents": reports},
        indent=1), encoding="utf-8")
    print(f"\nReport written to {out}")
    return reports


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(Path(__file__).parent / "samples"))
    run(Path(ap.parse_args().dir))
