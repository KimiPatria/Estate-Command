"""
Generate synthetic sample contract documents for the ingestion dry run.

We have no real contract files in the repo yet, so these synthesize the
structures the supply-chain use case will actually face — clause hierarchies,
force-majeure/mitigation language, commercial tables, a nested table, and a
two-column PDF layout — so dry_run.py has something representative to probe.
Replace with a handful of real (sanitised) contracts as soon as any are
available; findings from synthetic files are directional only.

Usage:  python doc_ingestion/generate_samples.py     # writes ./doc_ingestion/samples/
"""

from pathlib import Path

SAMPLES = Path(__file__).parent / "samples"

_FM_CLAUSE = (
    "Neither Party shall be liable for any failure or delay in performing its "
    "obligations under this Agreement to the extent that such failure or delay "
    "is caused by a Force Majeure Event, including but not limited to: acts of "
    "God, flood, drought, earthquake, volcanic haze, epidemic, export "
    "restriction, port closure, or governmental moratorium on palm oil "
    "shipments. The affected Party shall notify the other Party within seven "
    "(7) days and shall use commercially reasonable efforts to mitigate the "
    "effect of the Force Majeure Event."
)

_MITIGATION_CLAUSE = (
    "In the event of a supply disruption affecting Material delivery, Seller "
    "shall (a) allocate available stock pro-rata among its customers, (b) "
    "propose substitute materials of equivalent specification within fourteen "
    "(14) days, and (c) bear demurrage costs arising from re-routing, up to a "
    "cap of 2% of the affected Purchase Order value."
)


def make_docx_simple():
    from docx import Document
    doc = Document()
    doc.add_heading("SUPPLY AGREEMENT — CPO-2026-041", level=0)
    doc.add_paragraph("Between PT Wilmar Nabati Indonesia (\"Buyer\") and "
                      "PT Sumber Sawit Lestari (\"Seller\").")
    doc.add_heading("1. Scope of Supply", level=1)
    doc.add_paragraph("Seller shall supply Crude Palm Oil (Material Code "
                      "MAT-CPO-1001) meeting the specification in Schedule A.")
    doc.add_heading("2. Term", level=1)
    doc.add_paragraph("This Agreement is effective from 1 January 2026 to "
                      "31 December 2026 (the \"Active Period\").")
    doc.add_heading("3. Force Majeure", level=1)
    doc.add_paragraph(_FM_CLAUSE)
    doc.add_heading("4. Mitigation", level=1)
    doc.add_paragraph(_MITIGATION_CLAUSE)

    doc.add_heading("Schedule A — Commercial Terms", level=1)
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    header = ["Material Code", "Description", "Qty (MT)", "Price (USD/MT)"]
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = h
    data = [
        ["MAT-CPO-1001", "Crude Palm Oil", "5,000", "812.50"],
        ["MAT-PK-2044", "Palm Kernel", "1,200", "540.00"],
        ["MAT-PKO-3010", "Palm Kernel Oil", "600", "1,105.00"],
    ]
    for r, row in enumerate(data, start=1):
        for i, v in enumerate(row):
            table.rows[r].cells[i].text = v
    doc.save(SAMPLES / "contract_simple.docx")


def make_docx_nested():
    from docx import Document
    doc = Document()
    doc.add_heading("MASTER SERVICE AGREEMENT — LOG-2026-118", level=0)
    doc.add_paragraph("FFB logistics and evacuation services, Riau estates.")
    doc.add_heading("1. Service Levels", level=1)

    outer = doc.add_table(rows=2, cols=2)
    outer.style = "Table Grid"
    outer.rows[0].cells[0].text = "Region"
    outer.rows[0].cells[1].text = "Rate Card"
    outer.rows[1].cells[0].text = "Central Sumatra (K3)"
    # nested table inside a cell — the layout case Bedrock KB's default
    # parser is known to struggle with
    inner = outer.rows[1].cells[1].add_table(rows=3, cols=3)
    inner.style = "Table Grid"
    for i, h in enumerate(["Distance band", "Truck class", "IDR/tonne"]):
        inner.rows[0].cells[i].text = h
    inner.rows[1].cells[0].text = "0-25 km"
    inner.rows[1].cells[1].text = "6-wheel"
    inner.rows[1].cells[2].text = "85,000"
    inner.rows[2].cells[0].text = "25-60 km"
    inner.rows[2].cells[1].text = "10-wheel"
    inner.rows[2].cells[2].text = "132,000"

    doc.add_heading("2. Force Majeure", level=1)
    doc.add_paragraph(_FM_CLAUSE)
    doc.add_heading("3. Disruption Mitigation", level=1)
    doc.add_paragraph(_MITIGATION_CLAUSE)
    for n in range(4, 9):
        doc.add_heading(f"{n}. Boilerplate Section {n}", level=1)
        doc.add_paragraph("Standard terms " * 40)
    doc.save(SAMPLES / "contract_nested_tables.docx")


def make_pdf_two_column():
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "PURCHASE CONTRACT PC-2026-207 (TWO-COLUMN LAYOUT)",
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)

    left = ("ARTICLE 1 - MATERIAL. Buyer agrees to purchase fertiliser "
            "MAT-NPK-7734 (NPK 12-12-17) packed in 50kg bags. " * 6
            + "ARTICLE 2 - DELIVERY. DDP Dumai warehouse, partial shipments "
              "allowed. " * 6)
    right = ("ARTICLE 3 - FORCE MAJEURE. " + _FM_CLAUSE + " "
             "ARTICLE 4 - MITIGATION. " + _MITIGATION_CLAUSE)

    y0 = pdf.get_y()
    col_w = 90
    pdf.multi_cell(col_w, 4, left)
    pdf.set_xy(10 + col_w + 5, y0)
    pdf.multi_cell(col_w, 4, right)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 8, "Annex 1 - Delivery Schedule", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    rows = [["Lot", "Material", "Qty (bags)", "Window"],
            ["1", "MAT-NPK-7734", "20,000", "Feb 2026"],
            ["2", "MAT-NPK-7734", "20,000", "May 2026"],
            ["3", "MAT-MOP-5120", "8,000", "Aug 2026"]]
    for r in rows:
        for v in r:
            pdf.cell(45, 6, v, border=1)
        pdf.ln(6)
    pdf.output(str(SAMPLES / "contract_two_column.pdf"))


if __name__ == "__main__":
    SAMPLES.mkdir(exist_ok=True)
    make_docx_simple()
    make_docx_nested()
    make_pdf_two_column()
    print(f"Wrote 3 sample contracts to {SAMPLES}")
